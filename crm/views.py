import csv
import io
import json
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from portal.decorators import tool_access
from .forms import ContactForm, TagForm
from .config_forms import CRMConfigurationForm
from .models import Contact, ContactNote, Segment, Tag, CRMConfiguration
from .models import Attendance, StaffMember, StaffAttendance, InventoryTransaction
from .segment_forms import SegmentForm


@tool_access("crm")
def dashboard(request):
    contacts = Contact.objects.filter(owner=request.user)
    configuration = CRMConfiguration.objects.filter(owner=request.user).first()
    contact_groups = [{"value": value, "label": label, "contacts": contacts.filter(status=value), "count": contacts.filter(status=value).count()} for value, label in Contact.STATUS_CHOICES]
    return render(request, "crm/dashboard.html", {"total": contacts.count(), "leads": contacts.filter(status="lead").count(), "customers": contacts.filter(status="customer").count(), "inactive": contacts.filter(status="inactive").count(), "recent": contacts[:8], "contact_groups": contact_groups, "statuses": Contact.STATUS_CHOICES, "configuration": configuration})


@tool_access("crm")
def customize(request):
    configuration = CRMConfiguration.objects.filter(owner=request.user).first()
    form = CRMConfigurationForm(request.POST or None, instance=configuration)
    if request.method == "POST" and form.is_valid():
        saved = form.save(commit=False); saved.owner = request.user; saved.modules = form.cleaned_data["modules"]; saved.save(); messages.success(request, "CRM customization saved."); return redirect("crm:dashboard")
    return render(request, "crm/customize.html", {"form": form})


@tool_access("crm")
def contacts(request):
    qs = Contact.objects.filter(owner=request.user).prefetch_related("tags")
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "")
    if query:
        from django.db.models import Q
        qs = qs.filter(Q(name__icontains=query) | Q(company_name__icontains=query) | Q(email__icontains=query) | Q(phone__icontains=query))
    if status:
        qs = qs.filter(status=status)
    return render(request, "crm/contacts.html", {"contacts": qs, "query": query, "status": status, "statuses": Contact.STATUS_CHOICES})


@tool_access("crm")
def contact_create(request):
    form = ContactForm(request.POST or None, user=request.user)
    form.set_owner(request.user)
    if request.method == "POST" and form.is_valid():
        contact = form.save(commit=False); contact.owner = request.user; contact.save(); form.save_m2m(); messages.success(request, "Contact created."); return redirect("crm:contacts")
    return render(request, "crm/contact_form.html", {"form": form, "title": "Add contact"})


@tool_access("crm")
def contact_edit(request, contact_id):
    contact = get_object_or_404(Contact, id=contact_id, owner=request.user)
    form = ContactForm(request.POST or None, instance=contact, user=request.user)
    form.set_owner(request.user)
    if request.method == "POST" and form.is_valid(): form.save(); messages.success(request, "Contact updated."); return redirect("crm:contacts")
    return render(request, "crm/contact_form.html", {"form": form, "title": "Edit contact"})


@tool_access("crm")
def contact_delete(request, contact_id):
    contact = get_object_or_404(Contact, id=contact_id, owner=request.user)
    if request.method == "POST": contact.delete(); messages.success(request, "Contact deleted.")
    return redirect("crm:contacts")


@tool_access("crm")
def contact_status_update(request, contact_id):
    contact = get_object_or_404(Contact, id=contact_id, owner=request.user)
    if request.method == "POST":
        status = request.POST.get("status", "lead")
        if status in dict(Contact.STATUS_CHOICES):
            contact.status = status
            contact.save(update_fields=["status"])
            messages.success(request, "Contact status updated.")
    return redirect("crm:dashboard")


@tool_access("crm")
def contact_detail(request, contact_id):
    contact = get_object_or_404(Contact.objects.prefetch_related("tags", "activity_notes"), id=contact_id, owner=request.user)
    if request.method == "POST" and request.POST.get("note", "").strip():
        ContactNote.objects.create(contact=contact, author=request.user, text=request.POST["note"].strip())
        messages.success(request, "Note added.")
    return render(request, "crm/contact_detail.html", {"contact": contact})


@tool_access("crm")
def import_csv(request):
    if request.method == "POST" and request.FILES.get("file"):
        update_existing = request.POST.get("duplicate_mode") == "update"
        created = updated = skipped = 0
        uploaded = request.FILES["file"]
        if uploaded.name.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows)]
            source_rows = (dict(zip(headers, row)) for row in rows)
        else:
            source_rows = csv.DictReader(io.TextIOWrapper(uploaded.file, encoding="utf-8-sig"))
        for row in source_rows:
            phone, email = row.get("phone", "").strip(), row.get("email", "").strip().lower()
            existing = Contact.objects.filter(owner=request.user).filter(phone=phone).first() if phone else None
            existing = existing or (Contact.objects.filter(owner=request.user, email=email).first() if email else None)
            values = {"name": row.get("name", "").strip() or row.get("company_name", "").strip(), "company_name": row.get("company_name", "").strip(), "phone": phone, "email": email, "city": row.get("city", "").strip(), "category": row.get("category", "").strip(), "source": "csv", "status": row.get("status", "lead").strip() or "lead", "notes": row.get("notes", "").strip()}
            if existing and update_existing: Contact.objects.filter(pk=existing.pk).update(**values); updated += 1
            elif existing: skipped += 1
            else: Contact.objects.create(owner=request.user, **values); created += 1
        messages.success(request, f"Import complete. Created: {created}, updated: {updated}, skipped: {skipped}.")
        return redirect("crm:contacts")
    return render(request, "crm/import.html")


@tool_access("crm")
def tags(request):
    form = TagForm(request.POST or None)
    if request.method == "POST" and form.is_valid(): tag = form.save(commit=False); tag.owner = request.user; tag.save(); messages.success(request, "Tag created."); return redirect("crm:tags")
    return render(request, "crm/tags.html", {"form": form, "tags": Tag.objects.filter(owner=request.user)})


@tool_access("crm")
def export_json(request):
    data = list(Contact.objects.filter(owner=request.user).values("id", "name", "company_name", "phone", "email", "city", "category", "status", "source", "notes"))
    response = HttpResponse(json.dumps(data, indent=2), content_type="application/json")
    response["Content-Disposition"] = 'attachment; filename="crm_contacts.json"'; return response


@tool_access("crm")
def export_csv(request):
    response = HttpResponse(content_type="text/csv"); response["Content-Disposition"] = 'attachment; filename="crm_contacts.csv"'
    writer = csv.writer(response); writer.writerow(["name", "company_name", "phone", "email", "city", "category", "status", "source", "notes"])
    for c in Contact.objects.filter(owner=request.user): writer.writerow([c.name, c.company_name, c.phone, c.email, c.city, c.category, c.status, c.source, c.notes])
    return response


@tool_access("crm")
def segments(request):
    form = SegmentForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        segment = form.save(commit=False); segment.owner = request.user; segment.rules = {"field": form.cleaned_data["field"], "operator": form.cleaned_data["operator"], "value": form.cleaned_data["value"]}; segment.save(); messages.success(request, "Segment saved."); return redirect("crm:segments")
    return render(request, "crm/segments.html", {"form": form, "segments": Segment.objects.filter(owner=request.user)})


@tool_access("crm")
def bulk_action(request):
    if request.method == "POST":
        ids = request.POST.getlist("contact_ids"); qs = Contact.objects.filter(owner=request.user, id__in=ids); action = request.POST.get("action")
        if action == "delete": qs.delete()
        elif action == "status": qs.update(status=request.POST.get("status", "lead"))
        messages.success(request, f"Bulk action applied to {len(ids)} contacts.")
    return redirect("crm:contacts")


@tool_access("crm")
def attendance(request, contact_id):
    contact = get_object_or_404(Contact, id=contact_id, owner=request.user)
    if request.method == "POST": Attendance.objects.create(contact=contact, date=request.POST["date"], status=request.POST["status"], note=request.POST.get("note", "")); messages.success(request, "Attendance saved.")
    return render(request, "crm/attendance.html", {"contact": contact, "records": contact.attendance_records.all()})


@tool_access("crm")
def attendance_dashboard(request):
    from datetime import date
    selected_date = request.POST.get("date") or request.GET.get("date") or date.today().isoformat()
    date_from = request.GET.get("date_from") or selected_date
    date_to = request.GET.get("date_to") or selected_date
    members = list(StaffMember.objects.filter(owner=request.user, active=True).order_by("name"))
    selected_member = request.GET.get("member")
    if selected_member:
        members = [member for member in members if str(member.id) == selected_member]
    if request.method == "POST":
        for member in members:
            status = request.POST.get(f"status_{member.id}")
            if status:
                StaffAttendance.objects.update_or_create(member=member, date=selected_date, defaults={"status": status, "hours": request.POST.get(f"hours_{member.id}", member.standard_hours), "note": request.POST.get(f"note_{member.id}", "")})
        messages.success(request, "Attendance saved for the selected date.")
    records = {row.member_id: row for row in StaffAttendance.objects.filter(member__owner=request.user, date=selected_date)}
    range_rows = StaffAttendance.objects.filter(member__owner=request.user, date__range=[date_from, date_to]).select_related("member")
    present_range = range_rows.filter(status="present").count()
    absent_range = range_rows.filter(status="absent").count()
    late_range = range_rows.filter(status="late").count()
    salary = {}
    range_stats = {}
    for member in members:
        rows = [row for row in range_rows if row.member_id == member.id]
        present_days = sum(1 for row in rows if row.status == "present")
        absent_days = sum(1 for row in rows if row.status == "absent")
        late_days = sum(1 for row in rows if row.status == "late")
        total_hours = sum(row.hours for row in rows if row.status in ("present", "late"))
        salary[member.id] = total_hours * member.pay_rate if member.pay_type == "hourly" else (present_days + late_days) * member.pay_rate / 26
        range_stats[member.id] = {"present": present_days, "absent": absent_days, "late": late_days, "hours": total_hours}
    return render(request, "crm/attendance_dashboard.html", {"members": members, "records": records, "selected_date": selected_date, "date_from": date_from, "date_to": date_to, "present": sum(1 for row in records.values() if row.status == "present"), "absent": sum(1 for row in records.values() if row.status == "absent"), "late": sum(1 for row in records.values() if row.status == "late"), "present_range": present_range, "absent_range": absent_range, "late_range": late_range, "total": len(members), "salary": salary, "range_stats": range_stats, "range_rows": range_rows})


@tool_access("crm")
def add_attendance_member(request):
    if request.method == "POST":
        StaffMember.objects.create(owner=request.user, name=request.POST["name"], phone=request.POST.get("phone", ""), role=request.POST.get("role", ""), pay_type=request.POST.get("pay_type", "monthly"), pay_rate=request.POST.get("pay_rate", 0), standard_hours=request.POST.get("standard_hours", 8))
        messages.success(request, "Attendance member added.")
    return redirect("crm:attendance_dashboard") if request.method == "POST" else render(request, "crm/add_attendance_member.html")


@tool_access("crm")
def transactions(request):
    rows = list(InventoryTransaction.objects.filter(owner=request.user).select_related("contact"))
    month = request.GET.get("month", "")
    if month:
        rows = [row for row in rows if row.date.strftime("%Y-%m") == month]
    if request.method == "POST":
        contact_id = request.POST.get("contact") or None
        contact = get_object_or_404(Contact, id=contact_id, owner=request.user) if contact_id else None
        InventoryTransaction.objects.create(owner=request.user, kind=request.POST["kind"], item_name=request.POST["item_name"], quantity=request.POST["quantity"], unit_price=request.POST["unit_price"], date=request.POST["date"], contact=contact, note=request.POST.get("note", ""))
        messages.success(request, "Transaction saved."); return redirect("crm:transactions")
    sales = sum(row.total for row in rows if row.kind == "sale"); purchases = sum(row.total for row in rows if row.kind == "purchase")
    by_date = {}
    for row in rows:
        key = row.date.isoformat(); by_date.setdefault(key, {"date": key, "sales": 0, "purchases": 0}); by_date[key]["sales" if row.kind == "sale" else "purchases"] += float(row.total)
    return render(request, "crm/transactions.html", {"rows": rows, "contacts": Contact.objects.filter(owner=request.user), "sales": sales, "purchases": purchases, "net": sales - purchases, "month": month, "graph_data": list(sorted(by_date.values(), key=lambda item: item["date"]))})


@tool_access("crm")
def transaction_edit(request, transaction_id):
    transaction = get_object_or_404(InventoryTransaction, id=transaction_id, owner=request.user)
    if request.method == "POST":
        contact_id = request.POST.get("contact") or None
        contact = get_object_or_404(Contact, id=contact_id, owner=request.user) if contact_id else None
        transaction.kind = request.POST["kind"]; transaction.item_name = request.POST["item_name"]; transaction.quantity = request.POST["quantity"]; transaction.unit_price = request.POST["unit_price"]; transaction.date = request.POST["date"]; transaction.contact = contact; transaction.note = request.POST.get("note", ""); transaction.save(); messages.success(request, "Transaction updated."); return redirect("crm:transactions")
    return render(request, "crm/transaction_edit.html", {"transaction": transaction, "contacts": Contact.objects.filter(owner=request.user)})
